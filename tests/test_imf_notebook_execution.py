"""Offline execution of the shipped notebook's cells, not copied implementations."""

import ast
import csv
import json
import os
import sys
import time
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest

pd = pytest.importorskip("pandas")
pytest.importorskip("duckdb")
requests = pytest.importorskip("requests")
openpyxl = pytest.importorskip("openpyxl")
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill


NOTEBOOK = Path(__file__).parents[1] / "examples/notebooks/rlm_vs_plain_llm_imf_cpi.ipynb"
CELLS = json.loads(NOTEBOOK.read_text(encoding="utf-8"))["cells"]
# Whitespace-normalized snapshot fetched in memory from WEO_URL; SHA-256 matches
# cb0c2459c63b5d7d3d310fb6f0d6c32ee6624b79050dcf644e55a109a07bf360.
CORE_SENTENCE = (
    "Core inflation is expected to return to target only gradually in several major economies: "
    "by mid-2027 in the United Kingdom, by the end of 2027 in Japan and the United States, "
    "and only in 2028 in the euro area."
)


def source(marker):
    matches = ["".join(c["source"]) for c in CELLS if c["cell_type"] == "code"
               and marker in "".join(c["source"])]
    assert len(matches) == 1, marker
    return matches[0]


def execute(marker, namespace, definitions=False):
    tree = ast.parse(source(marker))
    if definitions:
        tree.body = [n for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef))]
    exec(compile(tree, f"{NOTEBOOK.name}:{marker}", "exec"), namespace)


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Network access is forbidden in notebook execution tests")

    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)
    monkeypatch.setattr("urllib.request.urlopen", forbidden)
    monkeypatch.setattr("socket.create_connection", forbidden)


@pytest.fixture
def namespace():
    ns = {}
    execute("def download_cpi", ns, definitions=True)
    execute("def grade_report", ns, definitions=True)
    execute("def streak_reference", ns, definitions=True)
    return ns


@pytest.fixture
def reference(tmp_path):
    path = tmp_path / "country's CPI.csv"
    header = ["COUNTRY", "INDEX_TYPE", "COICOP_1999", "TYPE_OF_TRANSFORMATION",
              "FREQUENCY", "TIME_PERIOD", "OBS_VALUE"]
    months = [f"{year}-M{month:02}" for year in range(2021, 2026) for month in range(1, 13)]
    records = []
    for country in range(18):
        for i, month in enumerate(months):
            records.append([f"C{country:02}", "CPI", "_T", "YOY_PCH_PA_PT", "M", month,
                            country * 10 + i // 12 + (i % 12) / 10])
    base = records[:60]
    # Missing, duplicated, nonnumeric, infinite and malformed months must not qualify.
    for name, observations in [("MISS", base[:-1]), ("DUP", base[:-1] + base[:1]),
                               ("EXTRA", base + base[:1])]:
        records.extend([[name, *row[1:]] for row in observations])
    for name, column, value in [("BAD", 6, "missing"), ("INF", 6, "inf"),
                                ("MONTH", 5, "2025-M13")]:
        observations = [[name, *row[1:]] for row in base]
        observations[-1][column] = value
        records.extend(observations)
    for column, value in [(1, "OTHER"), (2, "FOOD"), (3, "INDEX"), (4, "A")]:
        observations = [[f"FILTER{column}", *row[1:]] for row in base]
        for row in observations:
            row[column] = value
        records.extend(observations)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(reversed(records))
    # Closed-form expectations, independent of the notebook's SQL and streak algorithm.
    rows = [[f"C{c:02}", *[c * 10 + y + .55 for y in range(5)], c * 10 + 2.55]
            for c in reversed(range(18))]
    truth = {"n_countries": 18, "top10": rows[:10], "median_avg": 87.55,
             "all_countries": [[r[0], r[6]] for r in rows]}
    streaks = [[f"C{c:02}", 60, "2021-M01", "2025-M12", c * 10 + 5.1, "2025-M12"]
               for c in range(1, 18)]
    return path, truth, rows, streaks


def _pdf(*pages):
    document = MagicMock()
    document.__enter__.return_value = [SimpleNamespace(get_text=lambda page=page: page) for page in pages]
    return document


def report_workbook(path, truth, gap=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Average year-over-year CPI inflation (%), 2021-2025"])
    ws.merge_cells(f"A1:{'H' if gap is not None else 'G'}1")
    ws.append(["Country", 2021, 2022, 2023, 2024, 2025, "Avg 2021-2025"]
              + (["Gap vs world 2025 (pp)"] if gap is not None else []))
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    ws.column_dimensions["A"].width = 32
    for i, row in enumerate(truth["top10"]):
        ws.append(row + ([gap[i]] if gap is not None else []))
    ws["A13"] = "Median (all qualifying countries)"
    ws["G13"] = truth["median_avg"]
    allc = wb.create_sheet("All countries")
    allc.append(["Country", "Avg 2021-2025"])
    for cell in allc[1]:
        cell.font = Font(bold=True)
    for row in truth["all_countries"]:
        allc.append(row)
    for sheet in (ws, allc):
        for row in sheet.iter_rows(min_row=3 if sheet is ws else 2, min_col=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
    if gap is not None:
        wb.create_sheet("IMF Outlook")
    wb.save(path)
    return wb


def test_cells_compile_and_outputs_are_empty():
    for cell in CELLS:
        if cell["cell_type"] == "code":
            assert cell["outputs"] == []
            assert cell["execution_count"] is None
            code = "".join(cell["source"])
            if not code.startswith("%pip"):
                compile(code, cell["id"], "exec")


def test_ground_truth_and_grading_cells(namespace, reference, tmp_path):
    path, truth, rows, streaks = reference
    namespace["DATA_PATH"] = str(path)
    execute("con = duckdb.connect()", namespace)
    assert namespace["truth"] == truth
    assert namespace["streak_reference"](path) == (18, streaks)
    report = tmp_path / "report.xlsx"
    report_workbook(report, truth)
    namespace.update(REPORT_PATH=report, result=SimpleNamespace(submitted=True, payload={"n_countries": 18, "median_avg": 87.55}))
    execute("def grade_report", namespace)
    assert all(namespace["checks"].values())
    namespace["result"].submitted = False
    with pytest.raises(AssertionError):
        execute("def grade_report", namespace)
    namespace["result"].submitted = True
    namespace["result"].payload["n_countries"] = 17
    with pytest.raises(AssertionError):
        execute("def grade_report", namespace)
    namespace["con"].close()


@pytest.mark.parametrize("gap", [None, [1.25] * 10])
@pytest.mark.parametrize("mutation", [
    "sheet", "title", "merge", "header", "bold", "fill", "fill_pattern", "width",
    "format", "value", "country", "median_label", "median", "median_format",
    "all_header", "all_bold", "all_country", "all_value", "all_format", "all_order", "all_extra",
])
def test_report_rejects_wrong_workbooks(namespace, reference, tmp_path, gap, mutation):
    truth = reference[1]
    path = tmp_path / "report.xlsx"
    wb = report_workbook(path, truth, gap)
    assert all(namespace["grade_report"](path, truth, gap).values())
    ws, allc = wb["Report"], wb["All countries"]
    if mutation == "sheet":
        wb.create_sheet("Unexpected")
    elif mutation == "merge":
        ws.unmerge_cells(f"A1:{'H' if gap else 'G'}1")
    elif mutation in {"bold", "fill", "fill_pattern"}:
        if mutation == "bold":
            ws["G2"].font = Font(bold=False)
        else:
            ws["G2"].fill = PatternFill("solid" if mutation == "fill" else None, fgColor="FFFFFF" if mutation == "fill" else "D9D9D9")
    elif mutation == "width":
        ws.column_dimensions["A"].width = 31
    elif mutation in {"format", "median_format", "all_format"}:
        {"format": ws["G12"], "median_format": ws["G13"], "all_format": allc["B18"]}[mutation].number_format = "General"
    elif mutation == "all_bold":
        allc["B1"].font = Font(bold=False)
    elif mutation == "all_order":
        a, b = [c.value for c in allc[17]], [c.value for c in allc[18]]
        for col in (1, 2):
            allc.cell(17, col, b[col - 1])
            allc.cell(18, col, a[col - 1])
    elif mutation == "all_extra":
        allc.append(["XXX", 0])
    else:
        cell, value = {
            "title": (ws["A1"], "wrong"), "header": (ws["F2"], 2026),
            "value": (ws["F12"], ws["F12"].value + .01), "country": (ws["A12"], "XXX"),
            "median_label": (ws["A13"], "Median"), "median": (ws["G13"], 0),
            "all_header": (allc["B1"], "Average"), "all_country": (allc["A18"], "XXX"),
            "all_value": (allc["B18"], 1000),
        }[mutation]
        cell.value = value
    wb.save(path)
    assert not all(namespace["grade_report"](path, truth, gap).values()), mutation


@pytest.mark.parametrize("value", [None, "1.23", True, float("nan"), float("inf"), 1.24])
def test_numeric_checker_rejects_invalid_values(namespace, value):
    assert not namespace["ok"](value, 1.23)
    assert namespace["ok"](1.23000000001, 1.23)
    assert not namespace["ok_2dp"](value, 1.23)


def test_submitted_median_may_be_unrounded(namespace, reference, tmp_path):
    # One run submitted 4.683596270241045 for a true 4.68: the right figure, not rounded.
    assert namespace["ok_2dp"](1.2283596, 1.23) and namespace["ok_2dp"](1.23, 1.23)
    assert not namespace["ok_2dp"](1.2383596, 1.23) and not namespace["ok_2dp"](1.22, 1.23)
    path, truth, rows, streaks = reference
    report = tmp_path / "report.xlsx"
    report_workbook(report, truth)
    namespace.update(DATA_PATH=str(path), REPORT_PATH=report, truth=truth,
                     result=SimpleNamespace(submitted=True, payload={"n_countries": 18, "median_avg": 87.5496}))
    execute("def grade_report", namespace)
    assert all(namespace["checks"].values())
    namespace["result"].payload["median_avg"] = 87.56
    with pytest.raises(AssertionError):
        execute("def grade_report", namespace)


def test_tied_two_decimal_averages_may_follow_exact_values_or_codes(namespace):
    # GEO 5.8099 and DOM 5.8096 both show 5.81. Two of 24 real workbooks listed DOM first, by code, and
    # "sorted by the average, ties by country code, rounded to 2 decimals" allows that reading.
    same_order = namespace["same_order"]
    truth = [["LBN", 125.56], ["GEO", 5.81], ["DOM", 5.81], ["BBB", 4.46], ["CCC", 4.46], ["AAA", 4.46], ["ZZZ", 1.0]]
    by_exact = [tuple(row) for row in truth]
    by_code = [("LBN", 125.56), ("DOM", 5.81), ("GEO", 5.81), ("AAA", 4.46), ("BBB", 4.46), ("CCC", 4.46), ("ZZZ", 1.0)]
    assert same_order(by_exact, truth) and same_order(by_code, truth)
    neither = [("LBN", 125.56), ("GEO", 5.81), ("DOM", 5.81), ("CCC", 4.46), ("AAA", 4.46), ("BBB", 4.46), ("ZZZ", 1.0)]
    assert not same_order(neither, truth)
    assert not same_order([("GEO", 5.81), ("LBN", 125.56)] + by_exact[2:], truth)   # an untied pair swapped
    assert not same_order(by_exact[:-1], truth)                                      # a row short
    assert not same_order([("LBN", 125.5633)] + by_exact[1:], truth)                 # not rounded: one run wrote every cell this way
    assert not same_order([("LBN", "125.56")] + by_exact[1:], truth)                 # text, not a number
    assert not same_order([("XXX", 125.56)] + by_exact[1:], truth)                   # a country that does not belong


@pytest.mark.parametrize("damage", ["missing", "truncated", "partial_zip", "not_a_workbook"])
def test_unreadable_workbook_is_a_failed_check_not_a_traceback(namespace, reference, tmp_path, damage):
    # A save that raised partway left a 2 KB zip with no [Content_Types].xml; the grader died in zipfile.
    path = tmp_path / "report.xlsx"
    if damage == "partial_zip":
        import zipfile
        with zipfile.ZipFile(path, "w") as archive:  # what the failed save left behind
            archive.writestr("docProps/app.xml", "<Properties/>")
    elif damage != "missing":
        streak_workbook(path, 18, reference[3])
        path.write_bytes(path.read_bytes()[:900] if damage == "truncated" else b"not a workbook")
    result = SimpleNamespace(submitted=True, payload={"n_analyzed": 18, "n_with_streak": 17,
                                                     "longest_country": "C01", "longest_len": 60})
    checks = namespace["grade_streaks"](path, 18, reference[3], result)
    assert checks["workbook_opens"] is False and checks["submit_analyzed"]
    assert namespace["grade_report"](path, reference[1]) == {"workbook_opens": False}
    # A run that never submitted has no payload to read.
    unsubmitted = namespace["grade_streaks"](path, 18, reference[3], SimpleNamespace(submitted=False, payload=None))
    assert not any(unsubmitted.values())


def test_output_validator_sends_back_a_workbook_that_does_not_open(reference, tmp_path, monkeypatch):
    # The library accepts a payload when a validator raises anything but AssertionError,
    # so every way the file can be wrong has to surface as one.
    monkeypatch.setitem(sys.modules, "fabric_rlm", SimpleNamespace(RLM=None, File=None))
    ns = {}
    execute("rlm = RLM.task(", ns, definitions=True)
    path = tmp_path / "streaks.xlsx"
    validate = ns["workbook_opens"](path, ["Streaks", "Summary"])
    with pytest.raises(AssertionError, match="does not open"):
        validate({})
    streak_workbook(path, 18, reference[3])
    validate({"anything": 1})
    with pytest.raises(AssertionError, match="needs"):
        ns["workbook_opens"](path, ["Streaks", "Summary", "Missing"])({})
    path.write_bytes(path.read_bytes()[:900])
    with pytest.raises(AssertionError, match="does not open"):
        validate({})


def test_every_run_is_given_the_validator_for_its_own_workbook():
    for marker, path_name in [("rlm = RLM.task(", "REPORT_PATH"), ("rlm2 = RLM.task(", "REPORT2_PATH"),
                              ("result3 = rlm3.run()", "NOSKILLS_PATH")]:
        assert f"output_validator=workbook_opens({path_name}, [" in source(marker)
    two_source = source("rlm4 = RLM.task(")
    assert "output_validator=outlook_holds_together," in two_source
    assert 'workbook_opens(REPORT3_PATH, ["Report", "All countries", "IMF Outlook"])(payload)' in two_source


@pytest.mark.parametrize("marker,expected", [
    ("rlm = RLM.task(", {"n_countries": "int", "median_avg": "float"}),
    ("rlm2 = RLM.task(", {"n_analyzed": "int", "n_with_streak": "int", "longest_country": "str", "longest_len": "int"}),
    ("result3 = rlm3.run()", {"n_analyzed": "int", "n_with_streak": "int", "longest_country": "str", "longest_len": "int"}),
    ("rlm4 = RLM.task(", {"n_countries": "int", "median_avg": "float", "world_2025": "float", "world_2026": "float",
                          "revision_2026_pp": "float", "revision_2027_pp": "float", "core_target_economies": "int"}),
])
def test_every_run_declares_typed_outputs(marker, expected):
    # One run submitted the median as the text "4.68"; a list of names lets that through, a typed mapping sends it back.
    calls = [node for node in ast.walk(ast.parse(source(marker))) if isinstance(node, ast.Call)
             and isinstance(node.func, ast.Attribute) and node.func.attr == "task"]
    assert len(calls) == 1
    outputs = next(keyword.value for keyword in calls[0].keywords if keyword.arg == "outputs")
    assert isinstance(outputs, ast.Dict)
    assert {key.value: value.id for key, value in zip(outputs.keys, outputs.values)} == expected


@pytest.mark.parametrize("failure", ["stream", "http", "empty", "html", "header_only", "wrong_schema"])
@pytest.mark.parametrize("existing", [False, True])
def test_atomic_download_failure_and_retry(namespace, monkeypatch, tmp_path, failure, existing):
    path = tmp_path / "cpi.csv"
    old = b"previous file, possibly partial"
    if existing:
        path.write_bytes(old)
    header = b"COUNTRY,INDEX_TYPE,COICOP_1999,TYPE_OF_TRANSFORMATION,FREQUENCY,TIME_PERIOD,OBS_VALUE\n"
    valid = header + b"USA,CPI,_T,YOY_PCH_PA_PT,M,2021-M01,1.2\n"
    response = MagicMock()
    response.__enter__.return_value = response
    session = MagicMock()
    session.__enter__.return_value = session
    session.get.return_value = response
    monkeypatch.setattr(requests, "Session", lambda: session)

    def chunks():
        yield header
        assert path.read_bytes() == old if existing else not path.exists()
        raise requests.exceptions.ChunkedEncodingError("interrupted")

    response.iter_content.return_value = {
        "stream": chunks(), "http": [], "empty": [], "html": [b"<html>error</html>"],
        "header_only": [header], "wrong_schema": [b"a,b\n1,2\n"],
    }[failure]
    if failure == "http":
        response.raise_for_status.side_effect = requests.HTTPError("503")
    with pytest.raises((ValueError, requests.RequestException)):
        namespace["download_cpi"]("https://offline.invalid/cpi", path)
    assert path.read_bytes() == old if existing else not path.exists()
    response.__exit__.assert_called()
    response.raise_for_status.side_effect = None
    response.iter_content.return_value = [valid[:30], valid[30:]]
    namespace["download_cpi"]("https://offline.invalid/cpi", path)
    assert path.read_bytes() == valid
    assert not Path(str(path) + ".part").exists()
    # A second successful run must download again, not trust the existing header.
    response.iter_content.return_value = [valid.replace(b"1.2", b"2.3")]
    namespace["download_cpi"]("https://offline.invalid/cpi", path)
    assert path.read_bytes().endswith(b"2.3\n")
    assert session.get.call_count == 3
    assert session.get.call_args.kwargs["stream"] is True
    assert session.mount.call_args.args[1].max_retries.total == 3


def streak_workbook(path, analyzed, expected):
    wb = openpyxl.Workbook()
    st = wb.active
    st.title = "Streaks"
    st.append(["Longest high-inflation streaks (YoY >= 10%), 2021-2025"])
    st.merge_cells("A1:F1")
    st.append(["Country", "Months", "Start", "End", "Peak YoY", "Peak month"])
    for cell in st[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for col in ("A", "F"):
        st.column_dimensions[col].width = 14
    for row in expected[:15]:
        st.append(row)
    for row in range(3, 18):
        st.cell(row, 5).number_format = "0.00"
    st.conditional_formatting.add("E3:E17", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="FF0000"))
    chart = BarChart()
    chart.title = "Peak YoY, top 10"
    chart.add_data(Reference(st, min_col=5, min_row=3, max_row=12))
    chart.set_categories(Reference(st, min_col=1, min_row=3, max_row=12))
    st.add_chart(chart, "H2")
    summary = wb.create_sheet("Summary")
    for row in [["Countries analyzed", analyzed], ["Countries with a streak", len(expected)],
                ["Longest streak country", expected[0][0]], ["Longest streak months", expected[0][1]]]:
        summary.append(row)
    wb.save(path)
    return wb


STREAK_MUTATIONS_THAT_STILL_PASS = {None, "scale_repeated"}


@pytest.mark.parametrize("mutation", [None, "late_row", "summary", "format", "width", "fill", "chart", "anchor", "series", "categories", "chart_title", "scale",
                                      "scale_repeated", "scale_wrong_beside_right", "scale_missing", "scale_three_color", "note_row"])
def test_streak_grader(namespace, reference, tmp_path, mutation):
    path = tmp_path / "streaks.xlsx"
    expected = reference[3]
    wb = streak_workbook(path, 18, expected)
    st = wb["Streaks"]
    if mutation == "late_row":
        st["F17"] = "2021-M01"
    elif mutation == "summary":
        wb["Summary"]["B2"] = 1
    elif mutation == "format":
        st["E17"].number_format = "General"
    elif mutation == "width":
        st.column_dimensions["F"].width = 15
    elif mutation == "fill":
        st["F2"].fill = PatternFill("solid", fgColor="FFFFFF")
    elif mutation == "chart":
        st._charts.clear()
    elif mutation == "anchor":
        st._charts[0].anchor = "H3"
    elif mutation == "series":
        st._charts[0].series[0].val.numRef.f = "'Streaks'!$E$3:$E$11"
    elif mutation == "categories":
        st._charts[0].series[0].cat.numRef.f = "'Streaks'!$B$3:$B$12"
    elif mutation == "chart_title":
        st._charts[0].title = "Wrong"
    elif mutation == "scale":
        st.conditional_formatting["E3:E17"][0].colorScale.color[1].rgb = "0000FF"
    elif mutation == "scale_repeated":
        # Runs that could not read their rule back added the same rule again; Excel renders it as one.
        for _ in range(2):
            st.conditional_formatting.add("E3:E17", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="FF0000"))
    elif mutation == "scale_wrong_beside_right":
        st.conditional_formatting.add("E3:E17", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="00FF00"))
    elif mutation == "scale_missing":
        st.conditional_formatting._cf_rules.clear()
    elif mutation == "scale_three_color":
        st.conditional_formatting._cf_rules.clear()
        st.conditional_formatting.add("E3:E17", ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50,
                                                               mid_color="FFFF00", end_type="max", end_color="FF0000"))
    elif mutation == "note_row":
        st["A18"] = "Ranked by: streak length descending"
    wb.save(path)
    result = SimpleNamespace(submitted=True, payload={"n_analyzed": 18, "n_with_streak": 17,
                                                     "longest_country": "C01", "longest_len": 60})
    checks = namespace["grade_streaks"](path, 18, expected, result)
    assert all(checks.values()) == (mutation in STREAK_MUTATIONS_THAT_STILL_PASS), checks
    namespace.update(DATA_PATH=reference[0], REPORT2_PATH=path, result2=result)
    if mutation in STREAK_MUTATIONS_THAT_STILL_PASS:
        execute("def streak_reference", namespace)
    else:
        with pytest.raises(AssertionError):
            execute("def streak_reference", namespace)


@pytest.mark.parametrize("skills", [True, False])
@pytest.mark.parametrize("field", [None, "submitted", "n_analyzed", "n_with_streak", "longest_country", "longest_len"])
@pytest.mark.parametrize("missing", [False, True])
def test_streak_submission(namespace, reference, tmp_path, capsys, skills, field, missing):
    result = SimpleNamespace(submitted=True, payload={"n_analyzed": 18, "n_with_streak": 17,
                                                     "longest_country": "C01", "longest_len": 60},
                             n_turns=1, total_prompt_tokens=0, total_completion_tokens=0)
    if field == "submitted":
        result.submitted = False
    elif field:
        if missing:
            del result.payload[field]
        else:
            result.payload[field] = "WRONG" if field == "longest_country" else -1
    path = tmp_path / "streaks.xlsx"
    streak_workbook(path, 18, reference[3])
    namespace.update(DATA_PATH=reference[0], REPORT2_PATH=path, result2=result,
                     analyzed=18, expected_streaks=reference[3], DATA_DIR=str(tmp_path), STAMP="test",
                     Path=Path, os=os, time=time, STREAKS_TASK="offline", File=lambda p: p, lm_mini=None)

    def task(**kwargs):
        def run():
            streak_workbook(kwargs["inputs"]["report_path"], 18, reference[3])
            return result
        return SimpleNamespace(run=run)

    namespace["RLM"] = SimpleNamespace(task=task)
    if skills:
        if field is None:
            execute("def streak_reference", namespace)
        else:
            with pytest.raises(AssertionError):
                execute("def streak_reference", namespace)
        return
    # Without skills the cell is an experiment: it grades, prints both runs and never raises.
    good = SimpleNamespace(submitted=True, payload={"n_analyzed": 18, "n_with_streak": 17, "longest_country": "C01",
                                                    "longest_len": 60}, n_turns=3, total_prompt_tokens=10, total_completion_tokens=5)
    namespace.update(result2=good, streak_checks=namespace["grade_streaks"](path, 18, reference[3], good),
                     workbook_opens=lambda path, sheets: None)
    execute("result3 = rlm3.run()", namespace)
    shown = capsys.readouterr().out
    assert "with skills" in shown and "without skills" in shown
    failed = namespace["noskills_failed"]
    if field is None:
        assert failed == [] and "failed checks" not in shown
    else:
        expected = "submitted" if field == "submitted" else {"n_analyzed": "submit_analyzed", "n_with_streak": "submit_with_streak",
                    "longest_country": "submit_longest_country", "longest_len": "submit_longest_len"}[field]
        assert failed == [expected] and f"failed checks: {expected}" in shown


def test_report_paths_use_fresh_shared_stamp(namespace, tmp_path):
    # Execute actual setup assignments without downloads, Fabric clients or model calls.
    setup = ast.parse(source("def download_cpi"))
    setup.body = [n for n in setup.body if isinstance(n, (ast.Import, ast.ImportFrom))
                  or isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "STAMP" for t in n.targets)]
    namespace["DATA_DIR"] = str(tmp_path)
    paths = ["REPORT_PATH", "REPORT2_PATH", "NOSKILLS_PATH", "REPORT3_PATH"]
    previous = set()
    for _ in range(2):
        exec(compile(setup, "notebook-setup", "exec"), namespace)
        for name in paths:
            tree = ast.parse(source(f"{name} ="))
            tree.body = [n for n in tree.body if isinstance(n, ast.Assign)
                         and any(isinstance(t, ast.Name) and t.id == name for t in n.targets)]
            exec(compile(tree, "notebook-path", "exec"), namespace)
        current = {namespace[name] for name in paths}
        assert len(current) == 4
        assert not current & previous
        assert all(namespace["STAMP"] in Path(p).name for p in current)
        previous = current


@pytest.mark.parametrize("result_name,path_name", [
    ("result", "REPORT_PATH"), ("result2", "REPORT2_PATH"),
    ("result3", "NOSKILLS_PATH"), ("result4", "REPORT3_PATH"),
])
def test_failed_rerun_removes_stale_report(namespace, reference, tmp_path, monkeypatch, capsys, result_name, path_name):
    path = tmp_path / ("streaks_noskills_test.xlsx" if path_name == "NOSKILLS_PATH" else "report.xlsx")
    namespace.update(DATA_DIR=str(tmp_path), STAMP="test", DATA_PATH=reference[0], WEO_PATH="offline.pdf",
                     Path=Path, os=os, time=time, TASK="offline", STREAKS_TASK="offline",
                     TWO_SOURCE_TASK="offline", lm_mini=None, FabricLM=lambda *a, **kw: None,
                     workbook_opens=lambda path, sheets: None, fitz=SimpleNamespace(open=lambda path: _pdf("offline")))
    namespace[path_name] = str(path)

    def task(**kwargs):
        assert kwargs["inputs"]["report_path"] == str(path)
        assert not path.exists(), "Previous report survived into a new model call"
        return SimpleNamespace(run=lambda: SimpleNamespace(
            submitted=False, payload=None, report=lambda: "NOT SUBMITTED: ran out of turns", n_turns=12,
            total_prompt_tokens=0, total_completion_tokens=0))

    fake = SimpleNamespace(task=task)
    namespace.update(RLM=fake, File=lambda p: p)
    monkeypatch.setitem(sys.modules, "fabric_rlm", SimpleNamespace(RLM=fake, File=lambda p: p))
    tree = ast.parse(source(f"\n{result_name} = "))
    if result_name == "result3":
        # The run without skills does not stop the notebook: it reports that nothing was submitted.
        good = SimpleNamespace(submitted=True, payload={}, n_turns=3, total_prompt_tokens=0, total_completion_tokens=0)
        namespace.update(result2=good, streak_checks={"submitted": True}, analyzed=18, expected_streaks=reference[3])
        for _ in range(2):
            streak_workbook(path, 18, reference[3])
            exec(compile(tree, "notebook-failed-rerun", "exec"), namespace)
            assert not path.exists()
            assert namespace["noskills_failed"][0] == "submitted" and "workbook_opens" in namespace["noskills_failed"]
            assert "failed checks: submitted" in capsys.readouterr().out
        return
    # Run through the submission assertion, excluding subsequent display/grading.
    end = next(i for i, n in enumerate(tree.body) if isinstance(n, ast.Assert)
               and isinstance(n.test, ast.Attribute) and n.test.attr == "submitted")
    tree.body = tree.body[:end + 1]
    for _ in range(2):
        streak_workbook(path, 18, reference[3])
        with pytest.raises(AssertionError, match="NOT SUBMITTED: ran out of turns"):
            exec(compile(tree, "notebook-failed-rerun", "exec"), namespace)
        assert not path.exists()


def test_streak_reference_edges_and_ties(namespace, tmp_path):
    path = tmp_path / "edges.csv"
    months = [f"{y}-M{m:02}" for y in range(2021, 2026) for m in range(1, 13)]
    patterns = {"EARLY": [10, 12, 12] + [0] * 54 + [11, 12, 13],
                "END": [0] * 56 + [10, 11, 12, 12], "NONE": [9.99] * 60,
                "SAME": [10, 12, 12] + [0] * 57}
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["COUNTRY", "INDEX_TYPE", "COICOP_1999", "TYPE_OF_TRANSFORMATION", "FREQUENCY", "TIME_PERIOD", "OBS_VALUE"])
        for country, values in patterns.items():
            writer.writerows([country, "CPI", "_T", "YOY_PCH_PA_PT", "M", month, value]
                             for month, value in zip(months, values))
    assert namespace["streak_reference"](path) == (4, [
        ["END", 4, "2025-M09", "2025-M12", 12, "2025-M11"],
        ["EARLY", 3, "2021-M01", "2021-M03", 12, "2021-M02"],
        ["SAME", 3, "2021-M01", "2021-M03", 12, "2021-M02"],
    ])


TWO_SOURCE_MUTATIONS_THAT_STILL_PASS = {None, "labels_one_row_lower", "labels_no_empty_row", "median_unrounded",
                                        "sentence_without_period", "sentence_in_quotes", "driver_capitalized_with_period"}


@pytest.mark.parametrize("mutation", [None, "late_country", "gap", "gap_format", "label", "header", "bold", "format", "rate", "economy", "timing", "quote", "driver", "payload", "failed", "swap_uk_japan", "swap_uk_us", "year_only", "swap_quote",
                                      "labels_one_row_lower", "labels_no_empty_row", "median_unrounded", "median_wrong",
                                      "driver_label_missing", "labels_swapped", "source_label_not_bold", "sentence_not_below_label",
                                      "sentence_without_period", "sentence_in_quotes", "driver_capitalized_with_period",
                                      "sentence_cut_short", "driver_empty"])
def test_two_source_cell(namespace, reference, tmp_path, mutation):
    truth, rows = reference[1:3]
    path = tmp_path / "combined.xlsx"
    gap = [round(r[5] - 4.1, 2) for r in rows[:10]]
    wb = report_workbook(path, truth, gap)
    outlook = wb["IMF Outlook"]
    sentence = CORE_SENTENCE
    driver = "higher energy and food prices"
    for address, value in {
        "A1": "IMF World Economic Outlook Update, July 2026", "A2": "Measure", "B2": 2025, "C2": 2026, "D2": 2027,
        "A3": "World headline inflation (%)", "B3": 4.1, "C3": 4.7, "D3": 3.9,
        "A5": "Core inflation returns to target", "A6": "Economy", "B6": "Timing",
        "A7": "United Kingdom", "B7": "by mid-2027", "A8": "Japan", "B8": "by the end of 2027",
        "A9": "United States", "B9": "by the end of 2027", "A10": "euro area", "B10": "only in 2028",
        "A12": "Source sentence", "A13": sentence, "A15": "Stated driver of the 2026 increase", "A16": driver,
    }.items():
        outlook[address] = value
    for address in ["A2", "B2", "C2", "D2", "A5", "A6", "B6", "A12", "A15"]:
        outlook[address].font = Font(bold=True)
    for address in ["B3", "C3", "D3"]:
        outlook[address].number_format = "0.00"
    payload = {"n_countries": 18, "median_avg": 87.55, "world_2025": 4.1,
               "world_2026": 4.7, "revision_2026_pp": .3, "revision_2027_pp": .2, "core_target_economies": 4}

    def move_tail(shift):
        # Rows 12 to 16 hold the two labels and their texts; real runs left two empty rows above them, or none.
        tail = [(outlook.cell(row, 1).value, outlook.cell(row, 1).font.bold) for row in range(12, 17)]
        for row in range(11, 18):
            outlook.cell(row, 1).value = None
            outlook.cell(row, 1).font = Font(bold=False)
        for offset, (value, bold) in enumerate(tail):
            outlook.cell(12 + shift + offset, 1).value = value
            outlook.cell(12 + shift + offset, 1).font = Font(bold=bool(bold))

    if mutation == "labels_one_row_lower":
        move_tail(1)
    elif mutation == "labels_no_empty_row":
        move_tail(-1)
    elif mutation == "median_unrounded":
        payload["median_avg"] = 87.5496
    elif mutation == "median_wrong":
        payload["median_avg"] = 87.56
    elif mutation == "driver_label_missing":
        outlook["A15"] = None
    elif mutation == "labels_swapped":
        outlook["A12"], outlook["A15"] = outlook["A15"].value, outlook["A12"].value
    elif mutation == "source_label_not_bold":
        outlook["A12"].font = Font(bold=False)
    elif mutation == "sentence_not_below_label":
        outlook["A13"], outlook["B13"] = None, sentence
    elif mutation == "sentence_without_period":
        outlook["A13"] = sentence.rstrip(".")     # one real run split on the period and lost it
    elif mutation == "sentence_in_quotes":
        outlook["A13"] = "“" + sentence + "”"
    elif mutation == "driver_capitalized_with_period":
        outlook["A16"] = "Higher energy and food prices."
    elif mutation == "sentence_cut_short":
        outlook["A13"] = sentence.rsplit(",", 1)[0] + "."
    elif mutation == "driver_empty":
        outlook["A16"] = "."
    elif mutation == "late_country":
        wb["All countries"]["B18"] = 999
    elif mutation == "gap":
        wb["Report"]["H12"] = 999
    elif mutation == "gap_format":
        wb["Report"]["H12"].number_format = "General"
    elif mutation == "bold":
        outlook["A15"].font = Font(bold=False)
    elif mutation == "format":
        outlook["D3"].number_format = "General"
    elif mutation == "payload":
        payload["revision_2027_pp"] = 0
    elif mutation in {"swap_uk_japan", "swap_uk_us"}:
        other = "B8" if mutation == "swap_uk_japan" else "B9"
        outlook["B7"], outlook[other] = outlook[other].value, outlook["B7"].value
    elif mutation == "year_only":
        for address in ["B7", "B8", "B9"]:
            outlook[address] = "2027"
    elif mutation == "swap_quote":
        outlook["A13"] = sentence.replace("United Kingdom", "Japan").replace("Japan and", "United Kingdom and")
    elif mutation and mutation != "failed":
        address, value = {"label": ("A12", "Wrong"), "header": ("D2", 2028),
                          "rate": ("D3", 9), "economy": ("A10", "Mars"),
                          "timing": ("B10", "2029"), "quote": ("A13", "return to target only gradually"),
                          "driver": ("A16", "energy and food fabrication")}[mutation]
        outlook[address] = value
    wb.save(path)
    # Keep the actual pinned sentence at the PDF boundary, including when the workbook lies.
    document = MagicMock()
    document.__enter__.return_value = [SimpleNamespace(get_text=lambda: sentence + " The increase reflects " + driver + ".")]
    execute("rlm4 = RLM.task(", namespace, definitions=True)
    namespace.update(REPORT3_PATH=path, truth=truth, rows=rows, WEO_PATH="offline.pdf",
                     fitz=SimpleNamespace(open=lambda path: document),
                     result4=SimpleNamespace(submitted=mutation != "failed", payload=payload))
    if mutation in TWO_SOURCE_MUTATIONS_THAT_STILL_PASS:
        execute("WORLD_2025, WORLD_2026", namespace)
        assert all(namespace["two_source_checks"].values())
        assert len(namespace["core_rows"]) == 4
    else:
        with pytest.raises(AssertionError):
            execute("WORLD_2025, WORLD_2026", namespace)
        if mutation in {"swap_uk_japan", "swap_uk_us", "year_only"}:
            assert namespace["two_source_checks"]["source_sentence_quoted"]
            assert not namespace["two_source_checks"]["timings_correct"]


REPORT_PROSE = (
    # The real report uses the driver's words in more than one place, and the first is not the sentence about 2026:
    # a validator that looked at the first occurrence only sent back 17 of 30 correct workbooks.
    "Households in many economies faced higher energy and food prices last winter. "
    "Driven by surging energy prices, global headline inflation rose for a third month in a row year over year in May, "
    "breaking the downward trend that has been in place since the beginning of 2024. Headline inflation is projected "
    "to rise from 4.1 percent in 2025 to 4.7 percent in 2026 before easing to 3.9 percent in 2027, with the increase "
    "for 2026 driven mainly by higher energy and food prices. The forecast for 2026 is revised upward by 0.3 percentage "
    "point from the April 2026 WEO, whereas that for 2027 is revised upward by 0.2 percentage point. Growth in one "
    "economy is revised upward by 0.7 percentage point and reaches 10.4 percent. " + CORE_SENTENCE
    + " Risks are tilted to the downside."
)


def outlook_workbook(path, truth, **changes):
    """The two-source workbook as a correct run writes it, with named cells replaced."""
    wb = report_workbook(path, truth, [1.25] * 10)
    outlook = wb["IMF Outlook"]
    cells = {"A7": "United Kingdom", "B7": "by mid-2027", "A8": "Japan", "B8": "by the end of 2027",
             "A9": "the United States", "B9": "by the end of 2027", "A10": "euro area", "B10": "only in 2028",
             "A12": "Source sentence", "A13": CORE_SENTENCE, "A15": "Stated driver of the 2026 increase",
             "A16": "higher energy and food prices", "B3": 4.1, "C3": 4.7, "D3": 3.9}
    cells.update(changes)
    for address, value in cells.items():
        outlook[address] = value
    wb.save(path)


@pytest.fixture
def two_source_validator(reference, tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "fabric_rlm", SimpleNamespace(RLM=None, File=None))
    ns = {}
    execute("rlm = RLM.task(", ns, definitions=True)         # workbook_opens
    execute("rlm4 = RLM.task(", ns, definitions=True)        # bare, outlook_holds_together
    path = tmp_path / "combined.xlsx"
    ns.update(prose=REPORT_PROSE, REPORT3_PATH=path)
    return ns["outlook_holds_together"], path, reference[1]


@pytest.mark.parametrize("changes", [
    {},
    {"A13": CORE_SENTENCE.rstrip(".")},                                         # one real run lost the period
    {"A13": "\u201c" + CORE_SENTENCE + "\u201d"},
    {"A16": "Higher energy and food prices."},
    {"A16": "with the increase for 2026 driven mainly by higher energy and food prices"},
    {"A7": "the United Kingdom", "A10": "the euro area"},
    {"A12": None, "A13": "Source sentence", "A14": CORE_SENTENCE, "A15": None, "A16": "Stated driver of the 2026 increase",
     "A17": "higher energy and food prices"},                                  # labels one row lower, as two real runs wrote them
])
def test_two_source_validator_accepts_what_real_correct_runs_wrote(two_source_validator, changes):
    validate, path, truth = two_source_validator
    outlook_workbook(path, truth, **changes)
    validate({"core_target_economies": 4})
    validate({})                                                                # the count is checked only when submitted
    validate({"world_2025": 4.1, "world_2026": 4.7, "revision_2026_pp": 0.3, "revision_2027_pp": 0.2,
              "core_target_economies": 4, "n_countries": 18, "median_avg": 87.55})


@pytest.mark.parametrize("changes,payload,reason", [
    # The two slips real runs made, each sent to the grader as a finished report before this validator existed.
    ({"A10": "several major economies: by mid-2027 in the United Kingdom, by the end of 2027 in Japan and the United States, "
             "and only in 2028 in the euro area", "B10": "only gradually"}, {}, "not an economy's name"),
    ({"A16": "Driven by surging energy prices, global headline inflation rose for a third month in a row year over year in "
             "May, breaking the downward trend that has been in place since the beginning of 2024."}, {}, "not about 2026"),
    ({"A13": "Core inflation returns to target gradually in the major economies."}, {}, "not a sentence copied verbatim"),
    ({"A13": None}, {}, "not a sentence copied verbatim"),
    ({"B8": "in 2027"}, {}, "not worded as the source sentence words it"),
    ({"B9": None}, {}, "not worded as the source sentence words it"),
    ({"A9": "Canada"}, {}, "not an economy's name"),
    ({"A16": "rising energy and grocery bills"}, {}, "not in the report's own words"),
    ({"A16": None}, {}, "not in the report's own words"),
    ({"A12": "Source"}, {}, "no 'Source sentence' label"),
    ({"A15": None}, {}, "no 'Stated driver of the 2026 increase' label"),
    ({"A7": None}, {}, "No economy rows"),
    ({}, {"core_target_economies": 3}, "core_target_economies is 3 but the sheet has 4"),
    # One real run could not find the revision sentence, assumed 0.0 for both revisions and submitted.
    ({}, {"revision_2026_pp": 0.0, "revision_2027_pp": 0.0}, "revision_2026_pp is 0.0, but the report never says '0 percentage point'"),
    ({}, {"revision_2027_pp": 0.4}, "revision_2027_pp is 0.4"),
    ({}, {"revision_2026_pp": 4.1}, "revision_2026_pp is 4.1"),               # a rate is not a revision
    ({}, {"world_2025": 0.4}, "world_2025 is 0.4"),                            # "10.4 percent" must not count as 0.4
    ({}, {"world_2026": 0.3}, "world_2026 is 0.3"),                            # a revision is not a rate
    ({"D3": 9.9}, {}, "IMF Outlook D3 is 9.9"),
])
def test_two_source_validator_sends_back_a_slip_with_its_reason(two_source_validator, changes, payload, reason):
    validate, path, truth = two_source_validator
    outlook_workbook(path, truth, **changes)
    # AssertionError only: the library accepts the payload when a validator raises anything else.
    with pytest.raises(AssertionError, match=reason):
        validate(payload)


def test_two_source_validator_sends_back_a_workbook_that_does_not_open(two_source_validator):
    validate, path, truth = two_source_validator
    with pytest.raises(AssertionError, match="does not open"):
        validate({})
    outlook_workbook(path, truth)
    path.write_bytes(path.read_bytes()[:900])
    with pytest.raises(AssertionError, match="does not open"):
        validate({})


def test_two_source_scorecard_prints_before_the_cell_raises(namespace, reference, tmp_path, capsys):
    truth, rows = reference[1:3]
    path = tmp_path / "combined.xlsx"
    outlook_workbook(path, truth, A16="Driven by surging energy prices, global headline inflation rose in May.")
    wb = openpyxl.load_workbook(path)
    outlook = wb["IMF Outlook"]
    for address, value in {"A1": "IMF World Economic Outlook Update, July 2026", "A2": "Measure", "B2": 2025, "C2": 2026,
                           "D2": 2027, "A3": "World headline inflation (%)", "B3": 4.1, "C3": 4.7, "D3": 3.9,
                           "A5": "Core inflation returns to target", "A6": "Economy", "B6": "Timing"}.items():
        outlook[address] = value
    wb.save(path)
    execute("rlm4 = RLM.task(", namespace, definitions=True)
    payload = {"n_countries": 18, "median_avg": 87.55, "world_2025": 4.1, "world_2026": 4.7, "revision_2026_pp": .3,
               "revision_2027_pp": .2, "core_target_economies": 4}
    namespace.update(REPORT3_PATH=path, truth=truth, rows=rows, WEO_PATH="offline.pdf",
                     fitz=SimpleNamespace(open=lambda path: _pdf(REPORT_PROSE)),
                     result4=SimpleNamespace(submitted=True, payload=payload))
    with pytest.raises(AssertionError, match="failed these checks:.*driver_stated"):
        execute("WORLD_2025, WORLD_2026", namespace)
    shown = capsys.readouterr().out
    assert "FAIL  driver_stated" in shown and "PASS  timings_correct" in shown and "checks passed" in shown
