"""Grade RLM's self-authored workbook against the independent references.

Reads the workbook RLM wrote in Fabric (sheet "2. Answers") and compares each
answer with ground_truth.json, which was computed locally with pandas+DuckDB and
was NEVER uploaded to Fabric.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
_gt_list = json.loads((HERE.parent / "dbo_eval" / "ground_truth.json").read_text())
# q23 is withdrawn (reference unresolved) and was never sent to the agent.
GT = {r["id"]: r for r in _gt_list if r["id"] != "q23"}
WB = HERE / (sys.argv[1] if len(sys.argv) > 1 else "rlm_authored_full.xlsx")

TOL = 0.005  # 0.5% relative


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    m = re.search(r"-?\d+\.?\d*(?:[eE][-+]?\d+)?", s)
    return float(m.group()) if m else None


def close(a, b):
    if a is None or b is None:
        return False
    if b == 0:
        return abs(a) < 1e-9
    return abs(a - b) / abs(b) <= TOL


wb = openpyxl.load_workbook(WB)
ws = wb["2. Answers"]
hdr = [str(c.value or "").strip().lower() for c in ws[1]]
ci = {h: i for i, h in enumerate(hdr)}

rows = {}
order = []
dupes = []
for r in ws.iter_rows(min_row=2, values_only=True):
    qid = str(r[ci["question id"]] or "").strip().lower()
    if not qid:
        continue
    if qid in rows:
        dupes.append(qid)
    rows[qid] = r
    order.append(qid)

EXPECTED = len(GT)

print(f"workbook: {WB.name}  answers-sheet rows={ws.max_row - 1}  unique qids={len(rows)}")
print(f"order written: {order}")

# Structural assertions. A silent overwrite mid-run would otherwise disappear
# into the correctness percentage, so these are reported separately and first.
struct = {
    "unique_qids_equals_expected": len(rows) == EXPECTED,
    "no_duplicate_qids": not dupes,
    "no_missing_qids": set(rows) == set(GT),
}
print("\n--- structural checks (incremental-update integrity) ---")
for k, v in struct.items():
    print(f"  {'PASS' if v else 'FAIL'}  {k}")
if dupes:
    print(f"  !! duplicated/overwritten qids: {sorted(set(dupes))}")
missing_ids = sorted(set(GT) - set(rows))
if missing_ids:
    print(f"  !! qids never written to the workbook: {missing_ids}")

# Per-question published byte series, from the run log RLM's host wrote.
runlog = HERE / "run_log_cold_full.json"
if runlog.exists():
    log = json.loads(runlog.read_text())
    print("\n--- published workbook size after each question ---")
    prev = 0
    for rec in log:
        size = rec.get("published_bytes") or rec.get("bytes") or 0
        flag = "" if size > prev else "   <-- DID NOT GROW"
        print(f"  {rec.get('id','?'):5} {size:>7,} bytes  turns={rec.get('turns','?'):>3}{flag}")
        prev = max(prev, size)

results = []
for qid in sorted(GT):
    ref_val = GT[qid]["reference"]
    row = rows.get(qid)
    if row is None:
        results.append((qid, "MISSING", None, ref_val))
        continue
    raw = row[ci["answer"]]

    if isinstance(ref_val, str):
        # Categorical reference: the label must appear in the answer cell.
        text = str(raw or "").strip().lower()
        if not text:
            results.append((qid, "NO-ANSWER", None, ref_val))
            continue
        ok = ref_val.strip().lower() in text
        results.append((qid, "PASS" if ok else "FAIL", str(raw)[:40], ref_val))
        continue

    got = to_float(raw)
    if got is None:
        results.append((qid, "NO-ANSWER", str(raw)[:40] if raw else None, ref_val))
        continue
    results.append((qid, "PASS" if close(got, ref_val) else "FAIL", got, ref_val))

print(f"\n{'qid':6} {'verdict':10} {'rlm':>22} {'reference':>22}")
print("-" * 64)
for qid, verdict, got, ref in results:
    g = f"{got:,.4f}" if isinstance(got, float) else (str(got) if got is not None else "")
    r = f"{ref:,.4f}" if isinstance(ref, (int, float)) else str(ref)
    print(f"{qid:6} {verdict:10} {g[:22]:>22} {r[:22]:>22}")

n = len(results)
p = sum(1 for _, v, _, _ in results if v == "PASS")
miss = sum(1 for _, v, _, _ in results if v == "MISSING")
noans = sum(1 for _, v, _, _ in results if v == "NO-ANSWER")
present = [r for r in results if r[1] in ("PASS", "FAIL")]
print(f"\nrows lost to the workbook reset : {miss}/{n}")
print(f"present but not machine-readable: {noans}/{n}")
if present:
    print(f"correct among gradable rows     : {p}/{len(present)} = "
          f"{100*p/len(present):.1f}%")
print(f"correct over all {n} questions   : {p}/{n} = {100*p/n:.1f}%")
print("\nBaseline for reference only -- Phase 1 cold arm: 61/72 = 84.7%")
print("NOT a valid comparison: different model (gpt-4.1-mini), different")
print("      environment (local py3.11), and the HARNESS authored the workbook.")
print("      Compare arms within Phase 2 instead; see PHASE2_REPORT.md sec 2.")

(HERE / f"{WB.stem}_grade.json").write_text(json.dumps(
    {"structural": struct,
     "duplicates": sorted(set(dupes)),
     "missing": missing_ids,
     "correct": p, "total": n,
     "results": [{"qid": q, "verdict": v, "rlm": g, "reference": r}
                 for q, v, g, r in results]}, indent=2))
print(f"wrote {WB.stem}_grade.json")
