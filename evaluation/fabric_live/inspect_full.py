"""Download RLM's published workbook and inspect what it actually contains."""
import subprocess
import sys
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl

WS = "sandeep_ws"
LH = "da_agent_tests.Lakehouse"
REMOTE = "Files/rlm_excel_eval/out/cold-full-1/rlm_answers.xlsx"
LOCAL = Path(__file__).resolve().parent / "rlm_authored_full.xlsx"

tok = subprocess.run(
    ["az.cmd" if sys.platform == "win32" else "az", "account", "get-access-token",
     "--resource", "https://storage.azure.com", "--query", "accessToken", "-o", "tsv"],
    capture_output=True, text=True, check=True).stdout.strip()

url = f"https://onelake.dfs.fabric.microsoft.com/{WS}/{LH}/{REMOTE}"
req = Request(url, headers={"Authorization": f"Bearer {tok}",
                            "x-ms-version": "2023-08-03"}, method="GET")
with urlopen(req, timeout=300) as r:
    LOCAL.write_bytes(r.read())
print(f"downloaded {LOCAL.name}  ({LOCAL.stat().st_size:,} bytes)\n")

wb = openpyxl.load_workbook(LOCAL)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== {ws.title}   rows={ws.max_row}  cols={ws.max_column}")
    hdr = [c.value for c in ws[1]]
    print("   header:", hdr)
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 8), values_only=True):
        print("   ", [str(c)[:42] if c is not None else "" for c in row])

    # Did the model actually format it?
    h = ws.cell(1, 1)
    print(f"   header bold={h.font.bold} fill={h.fill.fgColor.rgb if h.fill and h.fill.fgColor else None} "
          f"freeze={ws.freeze_panes}")

