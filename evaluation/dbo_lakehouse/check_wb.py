import openpyxl

wb = openpyxl.load_workbook("dbo_eval_report.xlsx")
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== {ws.title}  dims={ws.dimensions}  rows={ws.max_row} cols={ws.max_column}")
    for r in ws.iter_rows(min_row=1, max_row=min(4, ws.max_row), values_only=True):
        print("   ", [str(c)[:34] if c is not None else "" for c in r][:9])

ws2 = wb["Questions"] if "Questions" in wb.sheetnames else wb.worksheets[1]
hdr = [str(c.value) for c in ws2[1]]
print("\nsheet2 header:", hdr)
qcol = 0
ids = set()
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row and row[qcol]:
        ids.add(str(row[qcol]))
print(f"sheet2 distinct question ids: {len(ids)} -> {sorted(ids)}")
