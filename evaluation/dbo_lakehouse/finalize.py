"""Post-run corrections to the deliverables.

1. Flag q23 (withdrawn) and q20 (three separate defects) in the workbook, so
   the workbook and the report no longer contradict each other.
2. Compute cold-vs-learned answer agreement -- the user's literal question
   was whether learned answers MATCH cold, which is distinct from whether
   each is correct.
"""
from __future__ import annotations

import json
import re
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill

WB = "dbo_eval_report.xlsx"

FLAGS = {
    "q23": ("WITHDRAWN -- defective reference. My SQL divides ALL usage_logs.api_calls "
            "by ACTIVE-user count, mismatching numerator and denominator populations. "
            "The library also disagreed with itself (3/6 on 10374.66), so this is "
            "excluded as UNRESOLVED, not as a clean library win. Not scored in the gate."),
    "q20": ("NOT a reference defect -- reference 20.8027 is correct. 0/6 for THREE "
            "unrelated reasons: (a) 4/6 trials used WHERE status='ACTIVE' uppercase "
            "against lowercase data -> empty result reported as a confident 0; "
            "(b) arm A rep1 returned 0.208 while declaring units='percentage' "
            "(missing x100); (c) arm A rep0 substituted subscription-count share "
            "for MRR share and reported status='ok'. Report separately."),
}

wb = openpyxl.load_workbook(WB)
ws = wb["2. Questions & Answers"]
hdr = [str(c.value) for c in ws[1]]
qcol = hdr.index("Q") + 1
ccol = hdr.index("Comments") + 1

amber = PatternFill("solid", fgColor="FFF2CC")
for row in range(2, ws.max_row + 1):
    qid = str(ws.cell(row, qcol).value or "")
    if qid in FLAGS:
        prev = str(ws.cell(row, ccol).value or "").strip()
        ws.cell(row, ccol).value = (prev + "  |  " if prev else "") + FLAGS[qid]
        for c in range(1, ws.max_column + 1):
            ws.cell(row, c).fill = amber
        ws.cell(row, qcol).font = Font(bold=True, color="9C5700")

wb.save(WB)
print(f"flagged {sorted(FLAGS)} in {WB}")

# ---------------------------------------------------------------- agreement
data = json.load(open("dbo_eval_results.json"))


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.search(r"[-+]?\d[\d,]*\.?\d*", v)
        if m:
            try:
                return float(m.group(0).replace(",", ""))
            except ValueError:
                return None
    return None


vals = defaultdict(dict)
for t in data["trials"]:
    if t.get("ok"):
        vals[(t["question_id"], t["rep"])][t["arm"]] = num(t.get("grade", {}).get("value"))

agree = disagree = unusable = 0
both_right = both_wrong = a_only = b_only = 0
gt = {r["id"]: r["reference"] for r in json.load(open("ground_truth.json"))}
by_q = defaultdict(lambda: [0, 0])

for (qid, rep), d in vals.items():
    if qid == "q23":
        continue
    a, b = d.get("A"), d.get("B")
    if a is None or b is None:
        unusable += 1
        continue
    ref = gt[qid]
    ok = (a != 0 and abs(b - a) <= abs(a) * 0.005) or (a == 0 and b == 0)
    agree += ok
    disagree += not ok
    by_q[qid][0] += ok
    by_q[qid][1] += 1
    if isinstance(ref, (int, float)):
        ra = abs(a - ref) <= abs(ref) * 0.005 if ref else a == 0
        rb = abs(b - ref) <= abs(ref) * 0.005 if ref else b == 0
        both_right += ra and rb
        both_wrong += (not ra) and (not rb)
        a_only += ra and not rb
        b_only += rb and not ra

tot = agree + disagree
print("\n" + "=" * 72)
print("COLD vs LEARNED -- do the answers MATCH?  (q23 excluded)")
print("=" * 72)
print(f"  comparable trial pairs      : {tot}")
print(f"  identical answer (<=0.5%)   : {agree}  ({100*agree/tot:.1f}%)")
print(f"  different answer            : {disagree}  ({100*disagree/tot:.1f}%)")
print(f"  pairs where one side had no parseable value: {unusable}")
print()
print("  Of comparable pairs, versus the reference:")
print(f"    both correct              : {both_right}")
print(f"    cold correct, learned NOT : {a_only}   <-- learning broke it")
print(f"    learned correct, cold NOT : {b_only}   <-- learning fixed it")
print(f"    both wrong                : {both_wrong}")
print()
print("  Questions where cold and learned NEVER agreed:")
for q in sorted(by_q):
    ok, n = by_q[q]
    if ok == 0:
        print(f"    {q}  (0/{n} reps agreed)")
